
import os
import uuid

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.dml.color import RGBColor

from config import Config


OUTPUT_FOLDER = Config.PPT_FOLDER
os.makedirs(OUTPUT_FOLDER, exist_ok=True)


# =====================================
# THEME COLOR DEFINITIONS
# =====================================

THEMES = {
    "Professional": {
        "primary": RGBColor(15, 23, 42),
        "secondary": RGBColor(217, 119, 6),
        "header_bg": RGBColor(15, 23, 42),
        "text_dark": RGBColor(30, 41, 59),
        "card_bg": RGBColor(248, 250, 252)
    },
    "Business": {
        "primary": RGBColor(30, 58, 138),
        "secondary": RGBColor(14, 165, 233),
        "header_bg": RGBColor(30, 58, 138),
        "text_dark": RGBColor(30, 41, 59),
        "card_bg": RGBColor(241, 245, 249)
    },
    "Modern": {
        "primary": RGBColor(79, 70, 229),
        "secondary": RGBColor(236, 72, 153),
        "header_bg": RGBColor(67, 56, 202),
        "text_dark": RGBColor(15, 23, 42),
        "card_bg": RGBColor(248, 250, 252)
    },
    "Technology": {
        "primary": RGBColor(15, 118, 110),
        "secondary": RGBColor(13, 148, 136),
        "header_bg": RGBColor(17, 94, 89),
        "text_dark": RGBColor(30, 41, 59),
        "card_bg": RGBColor(240, 253, 250)
    },
    "Education": {
        "primary": RGBColor(124, 45, 18),
        "secondary": RGBColor(234, 88, 12),
        "header_bg": RGBColor(154, 52, 18),
        "text_dark": RGBColor(30, 41, 59),
        "card_bg": RGBColor(255, 247, 237)
    }
}


def get_theme(theme_name):
    """Return the requested theme or Modern as the default."""
    if not theme_name:
        return THEMES["Modern"]

    for key, palette in THEMES.items():
        if key.lower() == theme_name.lower():
            return palette

    return THEMES["Modern"]


# =====================================
# PRESENTATION HELPER FUNCTIONS
# =====================================

def add_text(
    slide,
    text,
    left,
    top,
    width,
    height,
    font_size=18,
    color=None,
    bold=False,
    alignment=None
):
    """Add a consistently formatted text box."""
    box = slide.shapes.add_textbox(
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height)
    )

    text_frame = box.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    paragraph = text_frame.paragraphs[0]
    paragraph.text = str(text)
    paragraph.font.name = "Calibri"
    paragraph.font.size = Pt(font_size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color or RGBColor(30, 41, 59)

    if alignment is not None:
        paragraph.alignment = alignment

    return box


def add_speaker_notes(slide, notes):
    """Add presenter notes when notes are available."""
    if not notes:
        return

    try:
        slide.notes_slide.notes_text_frame.text = str(notes)
    except (AttributeError, ValueError):
        # Keep presentation generation working if notes are unavailable.
        pass


def normalize_points(points):
    """Normalize generated slide points into a list of strings."""
    if isinstance(points, str):
        points = [points]

    if not isinstance(points, list):
        return []

    return [
        str(point).strip()
        for point in points
        if point is not None and str(point).strip()
    ]


# =====================================
# CREATE PRESENTATION
# =====================================

def create_presentation(title, slides, theme_name="Modern"):
    """
    Build a widescreen PowerPoint from generated slide data.
    Supports detailed points and presenter notes.
    """
    prs = Presentation()

    # Widescreen 16:9 layout
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    palette = get_theme(theme_name)
    blank_layout = prs.slide_layouts[6]

    # Ensure valid input
    if not isinstance(slides, list):
        slides = []

    total_slides = len(slides)

    # -----------------------------------
    # 1. TITLE SLIDE
    # -----------------------------------

    title_slide = prs.slides.add_slide(blank_layout)

    top_bar = title_slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        0,
        0,
        prs.slide_width,
        Inches(0.4)
    )
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = palette["secondary"]
    top_bar.line.fill.background()

    card = title_slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(1.5),
        Inches(1.8),
        Inches(10.33),
        Inches(4.0)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = palette["card_bg"]
    card.line.color.rgb = palette["secondary"]
    card.line.width = Pt(1.5)

    add_text(
        title_slide,
        title,
        1.8, 2.2, 9.73, 1.8,
        font_size=36,
        color=palette["primary"],
        bold=True,
        alignment=PP_ALIGN.CENTER
    )

    add_text(
        title_slide,
        f"Executive Research Briefing | Theme: {theme_name} | AI-Assisted Synthesis",
        1.8, 4.2, 9.73, 1.0,
        font_size=16,
        color=palette["secondary"],
        bold=True,
        alignment=PP_ALIGN.CENTER
    )

    add_text(
        title_slide,
        "Powered by Groq High-Speed AI Engine",
        1.8, 4.95, 9.73, 0.5,
        font_size=12,
        color=RGBColor(148, 163, 184),
        alignment=PP_ALIGN.CENTER
    )

    # Optional notes for the title slide
    add_speaker_notes(
        title_slide,
        f"Introduce the presentation topic: {title}. "
        "Briefly explain the purpose and expected learning outcomes."
    )

    # -----------------------------------
    # 2. CONTENT SLIDES
    # -----------------------------------

    for i, item in enumerate(slides, start=1):

        if not isinstance(item, dict):
            continue

        slide_title = str(item.get("title", f"Slide {i}")).strip()

        # Avoid adding a duplicate title slide.
        if (
            i == 1
            and slide_title.lower() == str(title).strip().lower()
            and len(slides) > 1
        ):
            continue

        slide = prs.slides.add_slide(blank_layout)

        # -----------------------------------
        # HEADER BANNER
        # -----------------------------------

        header = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE,
            0,
            0,
            prs.slide_width,
            Inches(1.1)
        )
        header.fill.solid()
        header.fill.fore_color.rgb = palette["header_bg"]
        header.line.fill.background()

        stripe = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE,
            0,
            Inches(1.1),
            prs.slide_width,
            Inches(0.08)
        )
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = palette["secondary"]
        stripe.line.fill.background()

        # Slide title
        add_text(
            slide,
            slide_title,
            0.8, 0.16, 11.7, 0.65,
            font_size=26,
            color=RGBColor(255, 255, 255),
            bold=True
        )

        # Subtitle
        subtitle = str(item.get("subtitle", "")).strip()

        if subtitle:
            add_text(
                slide,
                subtitle.upper(),
                0.85, 0.82, 11.5, 0.22,
                font_size=9,
                color=palette["secondary"],
                bold=True
            )

        # -----------------------------------
        # CONTENT CARD
        # -----------------------------------

        body_card = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(0.8),
            Inches(1.5),
            Inches(11.73),
            Inches(5.3)
        )
        body_card.fill.solid()
        body_card.fill.fore_color.rgb = palette["card_bg"]
        body_card.line.color.rgb = RGBColor(226, 232, 240)
        body_card.line.width = Pt(1)

        # Content text box
        body_box = slide.shapes.add_textbox(
            Inches(1.2),
            Inches(1.8),
            Inches(10.95),
            Inches(4.7)
        )

        text_frame = body_box.text_frame
        text_frame.clear()
        text_frame.word_wrap = True
        text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

        text_frame.margin_left = Inches(0.08)
        text_frame.margin_right = Inches(0.08)
        text_frame.margin_top = Inches(0.08)
        text_frame.margin_bottom = Inches(0.08)

        points = normalize_points(item.get("points", []))

        if not points:
            points = ["No detailed content was generated for this slide."]

        # Adapt font size to the amount of content.
        if len(points) >= 7:
            font_size = 14
        elif len(points) >= 5:
            font_size = 16
        else:
            font_size = 18

        for idx, point in enumerate(points):
            paragraph = (
                text_frame.paragraphs[0]
                if idx == 0
                else text_frame.add_paragraph()
            )

            paragraph.text = f"•  {point}"
            paragraph.font.name = "Calibri"
            paragraph.font.size = Pt(font_size)
            paragraph.font.color.rgb = palette["text_dark"]
            paragraph.space_after = Pt(10)
            paragraph.line_spacing = 1.08

        # -----------------------------------
        # PRESENTER NOTES
        # -----------------------------------

        speaker_notes = item.get("speaker_notes", "")

        if isinstance(speaker_notes, (dict, list)):
            speaker_notes = str(speaker_notes)

        add_speaker_notes(slide, speaker_notes)

        # -----------------------------------
        # FOOTER / SLIDE NUMBER
        # -----------------------------------

        footer_box = slide.shapes.add_textbox(
            Inches(11.2),
            Inches(6.9),
            Inches(1.3),
            Inches(0.35)
        )

        footer_paragraph = footer_box.text_frame.paragraphs[0]
        footer_paragraph.text = f"{len(prs.slides)}"
        footer_paragraph.font.name = "Calibri"
        footer_paragraph.font.size = Pt(11)
        footer_paragraph.font.color.rgb = RGBColor(148, 163, 184)
        footer_paragraph.alignment = PP_ALIGN.RIGHT

    # -----------------------------------
    # SAVE PRESENTATION
    # -----------------------------------

    filename = f"Research_Presentation_{uuid.uuid4().hex[:8]}.pptx"
    filepath = os.path.join(OUTPUT_FOLDER, filename)

    prs.save(filepath)

    return filepath
